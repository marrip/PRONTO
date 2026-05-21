import configparser
import glob
import logging
import os
import pandas
import pptx

# get tumor mutational burden label
def get_tmb_string(val):
	level = "Not available"
	if(0 <= val <= 5):
		level = "Lav"
	elif(5 < val <= 20):
		level = "Intermediær"
	elif(val > 20):
		level = "Høy"
	return "{} mut/Mb; {}\n".format(val, level)

# use glob to find file in different folder structures
def glob_tsoppi_file(is_error, root, run_id, *path_units):
	glob_string_ous = os.path.join(root, '{}_TSO_500_LocalApp_postprocessing_results'.format(run_id), *path_units)
	glob_string_hus = os.path.join(root, run_id, '{}_ppr'.format(run_id), *path_units)
	files = []
	for gl in (glob_string_ous, glob_string_hus):
		files.extend(glob.glob(gl))
	if len(files) == 1:
		return files[0]
	elif not is_error:
		logging.warning("unsuccessful glob strings for {}:\n{}\n{}".format(run_id, glob_string_ous, glob_string_hus))
	else:
		logging.error("unsuccessful glob strings for {}:\n{}\n{}".format(run_id, glob_string_ous, glob_string_hus))
		raise ValueError

# normalize dataframe to expected column indices
def normalize_column_index(df: pandas.DataFrame, exp_col_idx: list):
	# determine current, missing and additional column indices
	curr_col_idx = df.columns.tolist()
	miss_col_idx = list(set(exp_col_idx) - set(curr_col_idx))
	add_col_idx = list(set(curr_col_idx) - set(exp_col_idx))
	# add missing column indices
	for i in miss_col_idx:
		df[i] = ' '
	# combine expected with additional to get all present column indices and rearrange columns accordingly, additional columns are moved to the right
	all_col_idx = exp_col_idx + add_col_idx
	return df[all_col_idx]

# set dataframe column format to 2 decimal points if float type
def set_column_to_2_decimals(df: pandas.DataFrame, col_name: str):
	if col_name in df.columns:
		if df[col_name].dtype == float:
			df.loc[:, col_name] = df[col_name].round(2).astype(str)
	else:
		logging.info("Column {} not found in dataframe".format(col_name))
	return df

# get data fitting on one slide based on slide index and max rows per slide
def get_slide_table_data(df: pandas.DataFrame, slide_idx: int, max_rows: int):
	start = slide_idx * max_rows
	stop = min(start + max_rows, len(df))
	if start >= len(df):
		return []
	table = df.values.tolist()
	header = [df.columns.tolist()]
	table_data = header + table[start:stop]
	return table_data

# add constructed table name to slide and format the textbox
def add_table_name(shapes: pptx.shapes.shapetree.SlideShapes, table_name: str, left: float, top: float, width: float, height: float, font_size: float, print_row_num: bool, slide_idx: int, total_slides: int, rows: int):

	# add textbox to slide
	paragraph = shapes.add_textbox(pptx.util.Inches(left), pptx.util.Inches(top), pptx.util.Inches(width), pptx.util.Inches(height)).text_frame.paragraphs[0]

	# construct table name with optional row number and slide count
	part_1 = ", Page {}/{}".format(slide_idx + 1, total_slides) if total_slides > 1 else ''
	part_2 = " (N={}{})".format(rows, part_1) if print_row_num else ''
	paragraph.text = "{}{}".format(table_name, part_2)

	# font formatting and placement
	paragraph.font.size = pptx.util.Pt(font_size)
	paragraph.font.bold = True
	paragraph.alignment = pptx.enum.text.PP_ALIGN.CENTER

# parse topfilter sections in config file and construct list of topfilter dictionaries
def parse_topfilter(cfg : configparser.ConfigParser, output_dir: str) -> list:
	top_filter_dict = []
	for section in cfg.sections():
		if section.startswith("FILTER0-"):
			filter = dict(cfg[section])
			filter["filter_columns"] = filter["filter_column"].split(",")
			if "filter_column_add" not in filter:
				filter["filter_column_add"] = None
			if "min_depth_tumor_dna" not in filter:
				filter["min_depth_tumor_dna"] = None
			filter["pre_table_output_path"] = "{}_{}_pre.txt".format(output_dir, filter["output_table"])
			filter["table_output_path"] = "{}_{}.txt".format(output_dir, filter["output_table"])
			top_filter_dict.append(filter)
	return top_filter_dict

# Filter small variant data based on keyword being present in filter column
def filter_small_variant_data(data: pandas.DataFrame, sample_id: str, filter_column: str, keyword: str) -> pandas.DataFrame:

	# check if required columns are present in data
	for column_name in [filter_column, "IGV_QC", "Class_judgement", "Sample_ID"]:
		if column_name not in data.columns:
			logging.error("Column {} not found in data".format(column_name))
			raise ValueError
	
	# only consider data for the specified sample_id
	data = data[data['Sample_ID'] == sample_id]

	# check if IGV_QC is "Not OK" but Class_judgement is not "exclude"
	if data[(data['IGV_QC'] == "Not OK") & (data['Class_judgement'] != "exclude")].shape[0] > 0:
		logging.error("""Dataset error: 
		IGV_QC is 'Not OK', but Class_judgement is not 'exclude'. Please check the QC Excel file and fix the mistake before run this script again!
		""")
		raise ValueError
	
	# filter according to keyword being present or not in filter column
	if keyword.startswith('!'):
		keys = keyword.replace('!', '').split(' && ')
		for key in keys:
			data = data[~data[filter_column].str.contains(key)] # remove rows that contain the key in filter column
	else:
		data = data[data[filter_column].str.contains(keyword.replace(',', '|'), na=False)] # only keep rows that contain the keyword in filter column

	return data
